from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from docx.oxml import OxmlElement

BASE = Path('Documentação/Documentacao_Sabor_e_Clic_3_Bimestre_Revisada.docx')
OUT = Path('Documentação/Documentacao_Sabor_e_Clic_3_Bimestre_V3_Base_Revisada.docx')
doc = Document(BASE)


def find_exact(text):
    return next(p for p in doc.paragraphs if p.text.strip() == text)


def find_starts(prefix):
    return next(p for p in doc.paragraphs if p.text.strip().startswith(prefix))


def has_page_break(p):
    return bool(p._p.xpath('.//w:br[@w:type="page"]'))


def replace_p(p, text):
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    pPr = p._p.pPr
    for child in list(p._p):
        if child is not pPr:
            p._p.remove(child)
    run = p.add_run(text)
    if rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, rpr)


def set_cell_font(cell, size=9):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if not trPr.xpath('./w:cantSplit'):
        trPr.append(OxmlElement('w:cantSplit'))


# Integração entre disciplinas
t0 = doc.tables[0]
t0.rows[2].cells[1].text = 'JavaScript modularizado por responsabilidade, DOM, carrinho, validações, cálculos em memória e perfis simulados.'
t0.rows[2].cells[2].text = 'Classes ES6+, Fetch/JSON, async/await e integração real com Flask no quarto bimestre.'
t0.rows[3].cells[1].text = 'Estrutura Flask inicial com controllers e endpoints simulados para autenticação, Reservas, Cardápio, Pedidos e KDS.'
t0.rows[3].cells[2].text = 'MVC completo, services, persistência, autenticação/RBAC e integração com o banco.'
for ri in (2, 3):
    for c in t0.rows[ri].cells:
        set_cell_font(c)

# KDS e arquitetura
replace_p(find_starts('No protótipo, os tickets são exemplos fixos'),
          'No protótipo, os tickets continuam sendo exemplos visuais, mas o back-end simulado já define uma consulta específica para o KDS. Nesta primeira entrega, a interface demonstra a fila e os estados sem sincronização real. No quarto bimestre, o KDS passará a consumir Pedidos persistidos pela API e enviará alterações de estado ao Flask, que validará as transições antes de atualizar o banco.')
replace_p(find_starts('O README registra uma estrutura futura modular'),
          'O repositório já separa parte das responsabilidades do front-end em arquivos JavaScript específicos e possui uma estrutura Flask inicial. No navegador, dados simulados sustentam as telas e os fluxos. No servidor, controllers exemplificam contratos HTTP sem persistência definitiva. A direção arquitetural continua sendo uma aplicação em camadas, na qual o front-end comunica-se apenas com Flask e o banco permanece acessível somente pelo back-end.')

# JavaScript
replace_p(find_starts('O objeto state concentra perfil'),
          'O JavaScript atual foi dividido por responsabilidade para evitar um arquivo monolítico e tornar explícitos os conteúdos do terceiro bimestre. data.js concentra arrays, objetos literais e estado compartilhado; site.js reúne navegação, localStorage e helpers; components.js gera componentes visuais; catalog.js manipula dinamicamente o DOM do Cardápio; cart.js mantém o carrinho e calcula valores; validation.js valida Reservas; login.js trata o login demonstrativo; views.js registra as telas; e a11y.js concentra comportamentos de acessibilidade. A navegação ocorre entre páginas HTML separadas, caracterizando uma MPA.')
replace_p(find_starts('function go(page){'),
          'function renderCatalog(category = "Todos") {\n  const grid = document.querySelector(".grid.cols-4");\n  grid.innerHTML = "";\n  products.forEach(product => {\n    if (category !== "Todos" && product.cat !== category) return;\n    const wrapper = document.createElement("div");\n    wrapper.innerHTML = productCard(product).trim();\n    grid.appendChild(wrapper.firstElementChild);\n  });\n}')
replace_p(find_starts('Trecho 3 - Navegação'), 'Trecho 3 - Renderização dinâmica do Cardápio utilizando forEach e manipulação do DOM.')
replace_p(find_starts('Utilizamos const, arrow functions'),
          'Utilizamos arrays, objetos literais, const, arrow functions, template literals, find, filter, reduce, forEach, for...of e for...in. O Cardápio é atualizado dinamicamente pelo DOM; o carrinho mantém itens e calcula subtotais e total em memória; os formulários possuem validações simples; e o login demonstrativo seleciona o perfil e armazena a sessão localmente. Esses elementos cobrem os conteúdos de JavaScript previstos para o terceiro bimestre.')
replace_p(find_starts('Na próxima etapa, pretendemos separar modelos'),
          'No quarto bimestre, a modularização será ampliada com classes ES6+, encapsulamento e módulos voltados à comunicação com a API. As chamadas Fetch serão assíncronas, utilizarão JSON e serão tratadas com async/await e try/catch. O estado que hoje é simulado no navegador passará a refletir as respostas reais do Flask.')

# Flask
replace_p(find_starts('As rotas a seguir representam a arquitetura'),
          'O back-end do terceiro bimestre é propositalmente simulado. Ele já possui ponto de inicialização, configuração da aplicação e controllers para autenticação, Reservas e Pedidos. Esses arquivos demonstram endpoints fundamentais e separação de responsabilidades, enquanto os models permanecem simples e sem persistência definitiva. A finalidade desta etapa é validar os contratos das operações antes da implementação completa do MVC, do banco e da integração com o front-end.')
replace_p(find_exact('14.1 Exemplo teórico de rota'), '14.1 Exemplo atual de rota simulada')
replace_p(find_starts('@bp.post("/api/sessoes/<int:id>/pedidos")'),
          '# POST /api/reservas - exemplo simplificado\ndef post_reserva(dados):\n    cliente = dados["cliente"]\n    bancada = dados["bancada"]\n    Reserva_model.post_reserva(cliente, bancada)\n    return True, 200')
replace_p(find_starts('Trecho 4 - Exemplo teórico'), 'Trecho 4 - Exemplo do back-end simulado para criação de Reserva.')
replace_p(find_starts('Nesse exemplo, a rota não calcula preços'),
          'O exemplo demonstra o contrato básico de uma operação: receber dados, encaminhá-los para a camada responsável e devolver um resultado com código HTTP. Nesta primeira entrega, os models são simulados. Na implementação definitiva, controllers e services aplicarão autenticação, validações, regras de negócio e transações antes de acessar o banco de dados.')
# remove o parágrafo de códigos HTTP para evitar uma página órfã
for p in list(doc.paragraphs):
    if p.text.strip().startswith('Utilizaremos 200 para leituras'):
        p._element.getparent().remove(p._element)
        break

# Endpoints simulados
t5 = doc.tables[5]
routes = [
    ('POST /api/cadastro', 'Cliente', 'Cadastro simulado.'),
    ('POST /api/login', 'Todos', 'Login simulado por perfil.'),
    ('GET /api/bancadas', 'Cliente/Admin', 'Lista Bancadas simuladas.'),
    ('POST /api/reservas', 'Cliente/Admin', 'Cria Reserva simulada.'),
    ('GET /api/reservas/{id}', 'Cliente/Admin', 'Consulta Reserva por ID.'),
    ('GET /api/cardapio', 'Cliente', 'Retorna Cardápio simulado.'),
    ('POST /api/pedido/cadastro', 'Cliente', 'Cria Pedido simulado.'),
    ('GET /api/pedido/{id}', 'Cliente/Admin', 'Consulta Pedido por ID.'),
    ('GET /api/kds/{id}', 'Cozinheiro/Admin', 'Retorna Pedido para o KDS.'),
]
for ri, vals in enumerate(routes, 1):
    for ci, val in enumerate(vals):
        t5.rows[ri].cells[ci].text = val
        set_cell_font(t5.rows[ri].cells[ci])
for row in t5.rows:
    cant_split(row)

# Estado do projeto
replace_p(find_starts('Concluímos o protótipo estático navegável'),
          'Concluímos o protótipo navegável, a identidade visual, as telas dos três atores, o Cardápio, o carrinho em memória, a visualização de Pedidos, o layout do KDS, o DER, o MER, o arquivo EER e o dicionário de dados. Também modularizamos o JavaScript para evidenciar arrays, objetos, for...of, for...in, forEach, DOM, validações e cálculos em memória, e estruturamos um back-end Flask simulado com endpoints fundamentais.')
replace_p(find_starts('Iniciamos a estrutura Python'),
          'Implementamos uma primeira estrutura Flask com configuração da aplicação e controllers simulados para cadastro/login, Bancadas/Reservas e Cardápio/Pedidos/KDS. Esses endpoints ainda não formam uma API persistente, mas já estabelecem contratos compatíveis com a evolução do sistema. A autenticação real, o banco físico e a comunicação Fetch/JSON permanecem para a próxima etapa.')
replace_p(find_starts('Nas próximas etapas, implementaremos a factory Flask'),
          'Nas próximas etapas, completaremos a factory Flask e a separação MVC, substituiremos os models simulados por persistência real, implementaremos autenticação e RBAC, criaremos a Sessão persistente, integraremos Cardápio, Pedido e KDS ao banco e conectaremos o front-end à API por Fetch/JSON. Depois serão acrescentados tratamento de loading/erros e testes para conflitos, permissões, totais e transições de estado.')

# Requisitos
t10 = doc.tables[10]
t10.rows[7].cells[1].text = 'Concluído para a etapa'
t10.rows[7].cells[2].text = 'data.js, components.js, catalog.js, cart.js e demais módulos: arrays, objetos, for...of, for...in, forEach, DOM e cálculos.'
t10.rows[8].cells[1].text = 'Concluído em nível inicial'
t10.rows[8].cells[2].text = 'validation.js e login.js demonstram validação no navegador; validação de autoridade será repetida no Flask.'
t10.rows[9].cells[1].text = 'Concluído como estrutura simulada'
t10.rows[9].cells[2].text = 'Aplicação Flask inicial, controllers e endpoints simulados de autenticação, Reservas, Cardápio, Pedidos e KDS.'
t10.rows[10].cells[1].text = 'Definido e demonstrado no protótipo'
t10.rows[10].cells[2].text = 'Perfis e separação visual existem; autenticação e RBAC de autoridade serão aplicados no servidor.'
for ri in (7, 8, 9, 10):
    for c in t10.rows[ri].cells:
        set_cell_font(c)
    cant_split(t10.rows[ri])

# Sumário
sum_refs = next(p for p in doc.paragraphs if p.text.strip() == 'Referências')
p = sum_refs.insert_paragraph_before('22 Estado atual da implementação e integração')
p.style = sum_refs.style

# Capítulo 22 antes das Referências, preservando a quebra original das Referências
refs = find_exact('REFERÊNCIAS')
refs_idx = next(i for i, p in enumerate(doc.paragraphs) if p._p is refs._p)
ref_break = doc.paragraphs[refs_idx - 1]
assert has_page_break(ref_break)
pb = ref_break.insert_paragraph_before()
pb.add_run().add_break(WD_BREAK.PAGE)
ref_break.insert_paragraph_before('22 ESTADO ATUAL DA IMPLEMENTAÇÃO E INTEGRAÇÃO', style='Heading 1')
ref_break.insert_paragraph_before('Esta seção consolida a implementação atual sem alterar a arquitetura definida anteriormente. No terceiro bimestre, o front-end e o Flask utilizam dados simulados para demonstrar os fluxos; no quarto bimestre, esses mesmos contratos serão ligados à persistência e à comunicação real por JSON.')
ref_break.insert_paragraph_before('22.1 Front-end e JavaScript', style='Heading 2')
ref_break.insert_paragraph_before('O front-end possui áreas separadas para Cliente, Cozinheiro e Administrador. O JavaScript está dividido entre data.js, site.js, components.js, catalog.js, cart.js, validation.js, login.js, views.js e a11y.js, cobrindo arrays, objetos, for...of, for...in, forEach, DOM, validação e cálculos em memória.')
ref_break.insert_paragraph_before('22.2 Back-end simulado', style='Heading 2')
ref_break.insert_paragraph_before('A estrutura Flask já demonstra cadastro/login, Bancadas/Reservas, Cardápio/Pedidos e KDS por meio de controllers e models simulados. A finalidade é validar endpoints e responsabilidades antes de conectar o banco e implementar autenticação e RBAC reais.')
ref_break.insert_paragraph_before('22.3 Integração e próxima etapa', style='Heading 2')
ref_break.insert_paragraph_before('O fluxo permanece Reserva → Sessão → Cardápio/Carrinho → Pedido → KDS. Na integração definitiva, o JavaScript enviará identificadores e quantidades ao Flask; o servidor validará a Sessão, recalculará valores, persistirá os dados e devolverá JSON. O KDS passará a consumir esses Pedidos e atualizar seus estados pelo back-end.')
for p in doc.paragraphs:
    if p.text.strip().startswith('22.'):
        p.paragraph_format.keep_with_next = True

doc.save(OUT)
print(OUT)
