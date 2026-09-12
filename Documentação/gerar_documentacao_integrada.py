from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
SRC = BASE / "Documentacao_Sabor_e_Clic_3_Bimestre_V2_Integrada.md"
OUT = BASE / "Documentacao_Sabor_e_Clic_3_Bimestre_V2_Integrada.docx"


def set_font(style, name="Arial", size=12, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    set_font(normal)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        style = doc.styles[style_name]
        set_font(style, size=size, bold=True)
        style.font.color.rgb = RGBColor(31, 78, 121)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    if "CodeBlock" not in doc.styles:
        code = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["CodeBlock"]
    set_font(code, "Consolas", 9)
    code.paragraph_format.left_indent = Cm(0.8)
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.line_spacing = 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Página ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_paragraph(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.7)
    p.add_run("• " + text.lstrip("• "))


def add_table(doc, rows):
    if not rows:
        return
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    table = doc.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.text = value
        shade(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for source_row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(source_row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5)
    doc.add_paragraph()


def parse_table_line(line):
    return [part.strip() for part in line.strip().strip("|").split("|")]


def build():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure_document(doc)

    i = 0
    first_content = True
    in_code = False
    code_lines = []

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph("\n".join(code_lines), style="CodeBlock")
                p.paragraph_format.first_line_indent = Cm(0)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue

        if not line:
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = [parse_table_line(x) for x in table_lines]
            parsed = [r for r in parsed if not all(set(c) <= {"-", ":", " "} for c in r)]
            add_table(doc, parsed)
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            title = line[2:]
            doc.add_heading(title, level=1)
        elif line.startswith("• ") or line.startswith("- "):
            add_bullet(doc, line[2:])
        elif first_content:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(line)
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(12)
        else:
            p = add_paragraph(doc, line)
            if line.endswith("— 2026") or line.startswith("Turma:") or line.startswith("Integrantes:") or line.startswith("Professores:"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)

        first_content = False
        i += 1

    doc.save(OUT)
    print(f"Gerado: {OUT}")


if __name__ == "__main__":
    build()
