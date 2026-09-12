from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

BASE = Path('Documentação/Documentacao_Sabor_e_Clic_3_Bimestre_Revisada.docx')
OUT = Path('Documentação/Documentacao_Sabor_e_Clic_3_Bimestre_V3_Base_Revisada.docx')

doc = Document(BASE)

# -------------------- helpers --------------------
def replace_terms_in_paragraph(p):
    repl = {'Party':'Sessão', 'PARTY':'SESSÃO', 'party':'sessão'}
    for run in p.runs:
        text = run.text
        for old, new in repl.items():
            text = text.replace(old, new)
        run.text = text

for p in doc.paragraphs:
    replace_terms_in_paragraph(p)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_terms_in_paragraph(p)

def replace_paragraph(startswith, new_text):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            p.clear()
            p.add_run(new_text)
            return True
    return False

def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet' if 'List Bullet' in doc.styles else 'Normal')
        p.add_run(item)

# -------------------- atualiza trechos já existentes --------------------
replace_paragraph(
    'Ao longo do terceiro bimestre, reunimos contribuições',
    'Ao longo do terceiro bimestre, reunimos contribuições de Artes Visuais e Design, JavaScript, Python, Banco de Dados e Segurança. O front-end evoluiu para uma aplicação multipágina navegável com telas específicas para Cliente, Cozinheiro e Administrador. Também modularizamos a lógica JavaScript e implementamos um back-end Flask simulado, com estrutura inicial da aplicação e endpoints fundamentais. A persistência definitiva e a comunicação real por Fetch/JSON permanecem para o quarto bimestre.'
)
replace_paragraph(
    'Nesta etapa concluímos o protótipo estático e a documentação conceitual de dados.',
    'Nesta etapa concluímos o protótipo navegável, a documentação conceitual de dados, as lógicas JavaScript exigidas no terceiro bimestre e uma estrutura inicial de back-end Flask com endpoints simulados. A aplicação ainda não persiste os dados do fluxo completo nem realiza integração real entre navegador e servidor; essa evolução será feita no quarto bimestre.'
)
replace_paragraph(
    'O README registra uma estrutura futura modular, com routes, controllers, models, services e middleware.',
    'O projeto adota separação por responsabilidades. No front-end, a lógica foi dividida em arquivos específicos para dados, navegação, componentes, cardápio, carrinho, validação, login, views e acessibilidade. No back-end simulado, run.py inicializa o Flask, config.py concentra configurações e os controllers exemplificam autenticação, reservas, cardápio, pedidos e KDS. A estrutura completa com routes, services, models e middleware permanece como evolução arquitetural.'
)
replace_paragraph(
    'O objeto state concentra perfil, página atual, carrinho e estado de Reserva.',
    'O JavaScript atual foi dividido por responsabilidade. data.js concentra arrays, objetos literais e estado compartilhado; site.js reúne navegação e helpers; components.js gera componentes visuais; catalog.js manipula dinamicamente o DOM do cardápio; cart.js mantém o carrinho e calcula totais; validation.js valida a reserva; login.js trata o login demonstrativo; views.js registra as telas; e a11y.js concentra comportamentos de acessibilidade. A navegação ocorre entre páginas HTML separadas, caracterizando uma MPA.'
)
replace_paragraph(
    'function go(page){',
    'function renderCatalog(category = "Todos") {\n  const grid = document.querySelector(".grid.cols-4");\n  grid.innerHTML = "";\n  products.forEach(product => {\n    if (category !== "Todos" && product.cat !== category) return;\n    const wrapper = document.createElement("div");\n    wrapper.innerHTML = productCard(product).trim();\n    grid.appendChild(wrapper.firstElementChild);\n  });\n}'
)
replace_paragraph(
    'Trecho 3 - Navegação e renderização dinâmica do protótipo.',
    'Trecho 3 - Renderização dinâmica do Cardápio utilizando forEach e manipulação do DOM.'
)
replace_paragraph(
    'Utilizamos const, arrow functions, template literals, map, filter, reduce, find e for...in.',
    'Utilizamos arrays, objetos literais, const, arrow functions, template literals, find, filter, reduce, forEach, for...of e for...in. O Cardápio é atualizado dinamicamente pelo DOM; o carrinho mantém itens e calcula subtotais e total em memória; os formulários possuem validações simples; e o login demonstrativo seleciona o perfil e armazena a sessão localmente. Esses elementos cobrem os conteúdos de JavaScript previstos para o terceiro bimestre.'
)
replace_paragraph(
    'Na próxima etapa, pretendemos separar modelos, API e páginas.',
    'No quarto bimestre, a modularização será ampliada com classes ES6+, encapsulamento e módulos voltados à comunicação com a API. As chamadas Fetch serão assíncronas, utilizarão JSON e serão tratadas com async/await e try/catch. O estado que hoje é simulado no navegador passará a refletir as respostas reais do Flask.'
)
replace_paragraph(
    'As rotas a seguir representam a arquitetura que definimos para completar o produto.',
    'O back-end do terceiro bimestre é propositalmente simulado. Ele já possui run.py, config.py, app/__init__.py e controllers de autenticação, reserva e pedidos. Esses arquivos demonstram os endpoints fundamentais e a separação de responsabilidades, mas os models ainda são stubs e a API não está integrada ao banco ou ao front-end. O objetivo desta etapa é validar o contrato das operações antes da implementação definitiva.'
)
replace_paragraph(
    '@bp.post("/api/sessoes/<int:id>/pedidos")',
    '# .../api/reservas - método POST\ndef post_reserva(dados):\n    cliente = dados["cliente"]\n    bancada = dados["bancada"]\n    Reserva_model.post_reserva(cliente, bancada)\n    return True, 200'
)
replace_paragraph(
    'Trecho 4 - Exemplo teórico da responsabilidade reduzida de uma rota Flask.',
    'Trecho 4 - Exemplo atual do back-end simulado para criação de Reserva.'
)
replace_paragraph(
    'Nesse exemplo, a rota não calcula preços nem consulta estoque diretamente.',
    'O exemplo atual demonstra o contrato básico de uma operação: receber dados, encaminhá-los para a camada responsável e devolver resultado e código HTTP. Nesta primeira entrega, os models são simulados. Na implementação definitiva, controllers e services aplicarão autenticação, validação, regras de negócio e transações antes de acessar o banco.'
)
replace_paragraph(
    'Concluímos o protótipo estático navegável, a identidade visual, as telas dos três atores',
    'Concluímos o protótipo navegável, a identidade visual, as telas dos três atores, o Cardápio, o carrinho em memória, a visualização de Pedidos, o layout do KDS, o DER, o MER, o arquivo EER e o dicionário de dados. Também modularizamos o JavaScript para evidenciar arrays, objetos, for...of, for...in, forEach, DOM, validações e cálculos em memória, e estruturamos um back-end Flask simulado com endpoints fundamentais.'
)
replace_paragraph(
    'Iniciamos a estrutura Python, a criação do banco e um exemplo de comunicação Fetch/JSON.',
    'Implementamos uma primeira estrutura Flask com configuração da aplicação e controllers simulados para cadastro/login, Bancadas/Reservas e Cardápio/Pedidos/KDS. Esses endpoints ainda não formam uma API persistente, mas já estabelecem os contratos que serão ligados ao front-end. A autenticação real, o banco físico e Fetch/JSON continuam como próxima etapa.'
)
replace_paragraph(
    'Nas próximas etapas, implementaremos a factory Flask, os módulos MVC, autenticação e RBAC',
    'Nas próximas etapas, completaremos a factory Flask e a separação MVC, substituiremos os models simulados por persistência real, implementaremos autenticação e RBAC, criaremos a Sessão persistente, integraremos Cardápio/Pedido/KDS ao banco e conectaremos o front-end à API por Fetch/JSON. Depois serão acrescentados tratamento de loading/erros e testes para conflitos, permissões, totais e transições de estado.'
)

# Atualiza tabela interdisciplinar existente.
if doc.tables:
    t = doc.tables[0]
    for row in t.rows[1:]:
        key = row.cells[0].text.strip()
        if key == 'JavaScript':
            row.cells[1].text = 'MPA navegável, módulos JS, arrays/objetos, loops modernos, DOM, validações, carrinho e cálculos em memória.'
            row.cells[2].text = 'Classes ES6+, módulos de API, Fetch/JSON, async/await e integração com Flask.'
        elif key == 'Python e Flask':
            row.cells[1].text = 'Estrutura Flask inicial e endpoints simulados para autenticação, Reservas, Cardápio, Pedidos e KDS.'
            row.cells[2].text = 'MVC completo, services/models reais, banco persistente, autenticação e rotas protegidas.'

# -------------------- nova seção complementar, sem remover imagens/conteúdo antigo --------------------
doc.add_page_break()
doc.add_heading('22 ESTADO ATUAL DA IMPLEMENTAÇÃO E INTEGRAÇÃO', level=1)
doc.add_paragraph(
    'Esta seção consolida o estado do repositório após a evolução do terceiro bimestre. Ela complementa os capítulos anteriores sem substituir a modelagem, as figuras, os protótipos e as decisões já documentadas.'
)

doc.add_heading('22.1 Páginas atuais do front-end', level=2)
add_bullets([
    'Acesso: index.html.',
    'Cliente: cliente_inicio.html, cliente_bancadas.html, cliente_reservas.html, cliente_sessao.html, cliente_cardapio.html, cliente_carrinho.html e cliente_pedidos.html.',
    'Cozinha: cozinha_dashboard.html, cozinha_kds.html, cozinha_pedidos.html, cozinha_estoque.html e cozinha_historico.html.',
    'Administrador: admin_dashboard.html, admin_reservas.html, admin_bancadas.html, admin_produtos.html, admin_ingredientes.html, admin_substituicoes.html, admin_usuarios.html e admin_relatorios.html.'
])

doc.add_heading('22.2 Módulos JavaScript atuais', level=2)
t = doc.add_table(rows=1, cols=3)
t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i,v in enumerate(['Arquivo','Responsabilidade','Requisito demonstrado']): t.rows[0].cells[i].text=v
for row in [
    ('data.js','Dados simulados e estado','arrays, objetos, for...in'),
    ('site.js','Navegação, helpers e sessão local','DOM, eventos e localStorage'),
    ('components.js','Componentes reutilizáveis','for...of e templates'),
    ('catalog.js','Cardápio dinâmico','forEach e manipulação do DOM'),
    ('cart.js','Carrinho e totais','find/filter/reduce/forEach e cálculos'),
    ('validation.js','Validação da Reserva','for...of e validação'),
    ('login.js','Login demonstrativo','validação e perfis'),
    ('views.js','Templates das páginas','renderização do protótipo'),
    ('a11y.js','Acessibilidade','comportamentos auxiliares')
]:
    cells=t.add_row().cells
    for i,v in enumerate(row): cells[i].text=v

doc.add_heading('22.3 Endpoints do back-end simulado', level=2)
t = doc.add_table(rows=1, cols=4)
t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i,v in enumerate(['Método','Endpoint','Finalidade','Estado']): t.rows[0].cells[i].text=v
for row in [
    ('POST','/api/cadastro','Cadastrar cliente','Simulado'),
    ('GET','/api/cadastro/<id>','Consultar usuário','Simulado'),
    ('POST','/api/login','Validar login','Simulado'),
    ('GET','/api/bancadas','Listar Bancadas','Simulado'),
    ('POST','/api/reservas','Criar Reserva','Simulado'),
    ('GET','/api/reservas/<id>','Consultar Reserva','Simulado'),
    ('GET','/api/cardapio','Listar Cardápio','Simulado'),
    ('POST','/api/pedido/cadastro','Criar Pedido vinculado à Reserva','Simulado'),
    ('GET','/api/pedido/<id>','Consultar Pedido','Simulado'),
    ('GET','/api/kds/<id>','Gerar consulta para o KDS','Simulado')
]:
    cells=t.add_row().cells
    for i,v in enumerate(row): cells[i].text=v

doc.add_heading('22.4 Integração conceitual entre os sistemas', level=2)
for text in [
    '1. O Cliente seleciona uma Bancada e envia os dados da Reserva.',
    '2. A Reserva cria o contexto que dará origem à Sessão.',
    '3. A Sessão ativa permitirá acesso ao Cardápio e criação de Pedidos.',
    '4. O JavaScript mantém o carrinho e apresenta os cálculos ao usuário.',
    '5. No quarto bimestre, o navegador enviará apenas identificadores e quantidades; o Flask recalculará preços e persistirá Pedido e itens.',
    '6. O KDS consumirá os Pedidos persistidos e enviará alterações de estado ao servidor.'
]: doc.add_paragraph(text)

doc.add_heading('22.5 Limite da simulação e evolução para o quarto bimestre', level=2)
doc.add_paragraph(
    'A filosofia adotada nesta entrega é demonstrar os conceitos estudados com exemplos simples e coerentes com a arquitetura final. O front-end usa mocks e localStorage; o Flask usa controllers e models simulados. O objetivo não é antecipar o sistema completo, mas evitar código descartável: os mesmos conceitos de Reserva, Sessão, Cardápio, Pedido e KDS serão mantidos quando a persistência, Fetch/JSON, autenticação e regras completas forem implementados.'
)

doc.add_heading('22.6 Situação dos requisitos principais do terceiro bimestre', level=2)
t = doc.add_table(rows=1, cols=3)
t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,v in enumerate(['Requisito','Situação','Evidência atual']): t.rows[0].cells[i].text=v
for row in [
    ('DER e dicionário de dados','Concluído','Arquivos e figuras já incorporados à documentação.'),
    ('Protótipo navegável','Concluído','Páginas de Cliente, Cozinha e Administrador.'),
    ('Cardápio Digital','Concluído como protótipo funcional','cliente_cardapio.html + catalog.js.'),
    ('KDS','Concluído como protótipo visual','cozinha_kds.html + views.js.'),
    ('Arrays e objetos','Concluído','data.js.'),
    ('for...of / for...in / forEach','Concluído','components.js, data.js, catalog.js e cart.js.'),
    ('DOM dinâmico','Concluído','catalog.js e scripts de interface.'),
    ('Validação de formulários','Concluído em nível inicial','validation.js e login.js.'),
    ('Cálculos em memória','Concluído','cart.js.'),
    ('Estrutura Flask','Concluída em nível inicial','run.py, config.py e app.'),
    ('Endpoints fundamentais','Concluídos como simulação','controllers Flask.'),
    ('Fetch/JSON integrado','Próxima etapa','Previsto para o quarto bimestre.')
]:
    cells=t.add_row().cells
    for i,v in enumerate(row): cells[i].text=v

doc.save(OUT)
print(OUT)
